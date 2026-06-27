import os
import logging
import threading
from datetime import datetime
import pytz

logger = logging.getLogger(__name__)
TZ = pytz.timezone(os.environ.get('TIMEZONE', 'Asia/Bangkok'))

# One lock per xlsx path so concurrent bills to the SAME group serialize
# (different groups still write in parallel). load → modify → save is not atomic.
_file_locks = {}
_locks_guard = threading.Lock()


def _lock_for(path: str) -> threading.Lock:
    with _locks_guard:
        lock = _file_locks.get(path)
        if lock is None:
            lock = threading.Lock()
            _file_locks[path] = lock
        return lock

THAI_MONTHS = [
    '', 'มกราคม', 'กุมภาพันธ์', 'มีนาคม', 'เมษายน', 'พฤษภาคม', 'มิถุนายน',
    'กรกฎาคม', 'สิงหาคม', 'กันยายน', 'ตุลาคม', 'พฤศจิกายน', 'ธันวาคม',
]

BILL_TYPE_MAP = {
    'receipt': 'ใบเสร็จรับเงิน',
    'invoice': 'ใบแจ้งหนี้',
    'quotation': 'ใบเสนอราคา',
}

COLUMNS = [
    'วันที่', 'เลขที่', 'ประเภท', 'ร้าน/บริษัท', 'ที่อยู่', 'เลขกำกับภาษี',
    'รายการ', 'ราคาก่อนภาษี', 'ส่วนลด', 'อัตราภาษี(%)', 'ภาษีมูลค่าเพิ่ม',
    'ยอดสุทธิ', 'วิธีชำระ', 'หมายเหตุ', 'ชื่อไฟล์', 'บันทึกเมื่อ',
]

COL_WIDTHS = [12, 16, 17, 25, 30, 20, 45, 16, 12, 13, 18, 15, 15, 20, 25, 20]


def _items_summary(items: list) -> str:
    if not items:
        return ''
    parts = []
    for item in items[:10]:
        name = item.get('name', '')
        qty = item.get('qty', 1)
        amount = item.get('amount', 0)
        parts.append(f'{name} x{qty} = {amount:,.2f}')
    return '; '.join(parts)


def _init_sheet(wb, sheet_name: str):
    """Create a sheet with styled header row and return it."""
    try:
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        ws = wb.create_sheet(title=sheet_name)
        ws.append(COLUMNS)
        return ws

    ws = wb.create_sheet(title=sheet_name)
    for col_idx, col_name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = Font(bold=True, color='FFFFFF', name='Sarabun')
        cell.fill = PatternFill(fill_type='solid', fgColor='7B2FBE')
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for i, w in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = 'A2'
    return ws


def write_bill_local(group_id: str, group_name: str, bill_data: dict,
                     source_file_path: str, storage_path: str) -> str:
    """
    Append bill row to local Excel:  STORAGE/group_id/bills/YEAR.xlsx
    Sheet name = Thai month (มกราคม … ธันวาคม).
    Returns the xlsx file path.
    """
    import openpyxl

    now = datetime.now(TZ)
    sheet_name = THAI_MONTHS[now.month]
    year = now.year

    bills_dir = os.path.join(storage_path, group_id, 'bills')
    os.makedirs(bills_dir, exist_ok=True)
    xlsx_path = os.path.join(bills_dir, f'{year}.xlsx')

    items_str = _items_summary(bill_data.get('items', []))
    row_data = [
        bill_data.get('date', ''),
        bill_data.get('receipt_no', ''),
        BILL_TYPE_MAP.get(bill_data.get('bill_type', ''), bill_data.get('bill_type', '')),
        bill_data.get('merchant', ''),
        bill_data.get('merchant_address', ''),
        bill_data.get('tax_id', ''),
        items_str,
        bill_data.get('subtotal') or 0,
        bill_data.get('discount') or 0,
        bill_data.get('tax_rate') or 7,
        bill_data.get('tax_amount') or 0,
        bill_data.get('total') or 0,
        bill_data.get('payment_method', ''),
        bill_data.get('note', ''),
        os.path.basename(source_file_path),
        now.strftime('%Y-%m-%d %H:%M:%S'),
    ]

    with _lock_for(xlsx_path):
        if os.path.isfile(xlsx_path):
            try:
                wb = openpyxl.load_workbook(xlsx_path)
            except Exception as e:
                # Corrupt/half-written file — back it up and start fresh so we
                # never lose new bills to an unreadable workbook.
                logger.error(f'Excel load failed ({e}); recreating {xlsx_path}')
                try:
                    os.rename(xlsx_path, xlsx_path + '.corrupt')
                except Exception:
                    pass
                wb = openpyxl.Workbook()
                if 'Sheet' in wb.sheetnames:
                    del wb['Sheet']
        else:
            wb = openpyxl.Workbook()
            if 'Sheet' in wb.sheetnames:
                del wb['Sheet']

        if sheet_name not in wb.sheetnames:
            _init_sheet(wb, sheet_name)

        wb[sheet_name].append(row_data)

        # Atomic save: write to temp then replace, so a crash mid-save can't
        # leave a truncated xlsx.
        tmp_path = xlsx_path + '.tmp'
        wb.save(tmp_path)
        os.replace(tmp_path, xlsx_path)

    logger.info(f'Bill written to {xlsx_path} [{sheet_name}]')
    return xlsx_path


def write_bill_async(group_id: str, group_name: str, bill_data: dict,
                     source_file_path: str, storage_path: str):
    """Write bill to local Excel then sync to GDrive — all in background."""
    def _worker():
        try:
            xlsx_path = write_bill_local(group_id, group_name, bill_data,
                                         source_file_path, storage_path)
            _sync_to_gdrive(group_id, group_name, xlsx_path)
        except Exception as e:
            logger.error(f'excel_writer worker failed: {e}')

    threading.Thread(target=_worker, daemon=True).start()


def _sync_to_gdrive(group_id: str, group_name: str, xlsx_path: str):
    """Upload/replace the Excel file in mochi/GROUP/bills/ on GDrive."""
    try:
        from handlers.gdrive import _get_service, get_group_folder_id, _get_or_create_folder, upload_file_to_folder
        service = _get_service()
        if not service:
            return
        group_folder_id = get_group_folder_id(service, group_id, group_name)
        bills_folder_id = _get_or_create_folder(service, 'bills', group_folder_id)
        upload_file_to_folder(service, xlsx_path, bills_folder_id, os.path.basename(xlsx_path))
        logger.info(f'Excel synced to GDrive: {os.path.basename(xlsx_path)}')
    except Exception as e:
        logger.error(f'Excel GDrive sync failed: {e}')
